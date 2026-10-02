"""Build the next 14 days as a calendar (.ics) from schedule.json, the job tracker and jobs.json,
then publish it to a secret GitHub Gist that Lyfe Time reads.

Placeholders in schedule titles:
  {outreach}    connection requests + follow-ups due, from the tracker's Contacts tab
  {apply:N}     N real open jobs from jobs.json not yet in the tracker's Applications tab
  {leetcode:N}  N LeetCode problems, rotating through the 5 core patterns
  {catchup}     whatever is still missed this week (Sunday)

Missed work carries over: Lyfe Time writes "done=k/N at=HH:MM" onto its [daily-plan] busy copies in
the primary Google calendar, which plan.py reads through the private iCal address. Counts missed earlier
in the week are added to the next block of the same type that week; what is left goes to {catchup}.

Usage:
  uv run plan.py            # write out/plan.ics and the tracker's Metrics tab
  uv run plan.py --publish  # also push the plan to the gist in schedule.json
"""
import argparse
import json
import re
import subprocess
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).parent
OUT = ROOT / "out"
DAYS = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]
PATTERNS = ["Two pointers / sliding window", "Hash map / set", "BFS / DFS", "Fast & slow pointers", "Binary search"]

FIXED = ("Work block", "Sleep", "Wind down", "Lunch", "Meal", "Wake up")  # real events never move these
DROP = ("Free",)  # free time simply gives way to real events
SOFT = ("Dinner prep", "Workspace setup", "Post-workout", "Gentle recovery", "Free")  # tasks may take their time
DAY_START, DAY_END = 7 * 60, 21 * 60 + 30  # tasks stay out of wind-down and sleep
# words that mark an event you can follow online while doing something else
ONLINE = ("http", "zoom", "meet.google", "teams.microsoft", "youtube", "webinar", "livestream", "live with",
          "live stream", "online", "virtual", "watch")
# countable task types: SUMMARY starts with the count
KINDS = {"apply": r"^(\d+) applications?\b", "leetcode": r"^(\d+) LeetCode\b",
         "outreach": r"^(\d+) connection requests\b", "referral": r"^(\d+) referral asks\b"}
NAMES = {"apply": "applications", "leetcode": "LeetCode", "outreach": "connection requests", "referral": "referral asks"}


def mins(hhmm):
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def hhmm(m):
    return f"{m // 60:02d}:{m % 60:02d}"


def span(s, e):
    """Block minutes; a block that ends at or before it starts (sleep) runs to midnight."""
    a, b = mins(s), mins(e)
    return a, (b if b > a else 24 * 60)


def kind_of(summary):
    for kind, rx in KINDS.items():
        m = re.match(rx, summary)
        if m:
            return kind, int(m.group(1))
    return None, 1


def rows(wb, name):
    ws = wb[name]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]


def as_date(v):
    return v.date() if isinstance(v, datetime) else v if isinstance(v, date) else None


def esc(text):
    return text.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


# ---------- Google calendar: real events and Lyfe Time's done markers ----------

def read_calendar(cfg, lo_day, hi_day):
    """Returns (in_person, online, progress).

    in_person / online: {date: [(start_min, end_min, title)]}. In-person events block time; online ones
    (webinars, livestreams, video calls) can run alongside a task, so they only add a note.
    progress: [{day, summary, kind, n, done, at, bed, wake}] from Lyfe Time's [daily-plan] copies.
    """
    urls = cfg.get("busy_ics") or []
    if not urls:
        return {}, {}, []
    import urllib.request
    from zoneinfo import ZoneInfo

    import icalendar
    import recurring_ical_events

    tz = ZoneInfo(cfg["tz"])
    lo = datetime.combine(lo_day, datetime.min.time(), tz)
    hi = datetime.combine(hi_day, datetime.min.time(), tz)
    busy, online, progress = {}, {}, {}
    for url in urls:
        with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "job-assistant"}), timeout=30) as r:
            cal = icalendar.Calendar.from_ical(r.read())
        for ev in recurring_ical_events.of(cal).between(lo, hi):
            desc = str(ev.get("DESCRIPTION", ""))
            key = re.search(r"\[daily-plan\] key=(\d{4}-\d{2}-\d{2}) (.+)", desc)
            if key:  # Lyfe Time's busy copy of a plan block: read progress, never treat as a real event
                done = re.search(r"done=(\d+)/(\d+)", desc)
                if done:
                    summary = key.group(2).strip()
                    kind, _ = kind_of(summary)
                    at, bed, wake = (re.search(rf"{f}=(\d{{2}}:\d{{2}})", desc) for f in ("at", "bed", "wake"))
                    progress[(key.group(1), summary)] = {
                        "day": date.fromisoformat(key.group(1)), "summary": summary, "kind": kind,
                        "n": int(done.group(2)), "done": int(done.group(1)),
                        "at": at and at.group(1), "bed": bed and bed.group(1), "wake": wake and wake.group(1)}
                continue
            s, e = ev["DTSTART"].dt, ev.get("DTEND", ev["DTSTART"]).dt
            if not isinstance(s, datetime) or str(ev.get("TRANSP", "OPAQUE")) == "TRANSPARENT":
                continue  # all-day and "free" events don't block time
            text = " ".join(str(ev.get(k, "")) for k in ("SUMMARY", "LOCATION", "DESCRIPTION")).lower()
            target = online if any(w in text for w in ONLINE) else busy
            s, e = s.astimezone(tz), e.astimezone(tz)
            day = s.date()
            while day <= e.date():  # split events that cross midnight
                a = mins(s.strftime("%H:%M")) if day == s.date() else 0
                b = mins(e.strftime("%H:%M")) if day == e.date() else 24 * 60
                if b > a:
                    target.setdefault(day, []).append((a, b, str(ev.get("SUMMARY", "Busy"))))
                day += timedelta(days=1)
    return busy, online, list(progress.values())


def missed_this_week(progress, today):
    """Units missed from Monday through today, per kind. Days he did not review count as unknown."""
    monday = today - timedelta(days=today.weekday())
    missed = defaultdict(int)
    for p in progress:
        if p["kind"] and monday <= p["day"] <= today:
            missed[p["kind"]] += max(0, p["n"] - p["done"])
    return missed


# ---------- tasks ----------

class Planner:
    def __init__(self, cfg, carry, today):
        wb = load_workbook(cfg["tracker"], data_only=True)
        self.contacts = rows(wb, "Contacts")
        applied = {(r["Company"] or "").lower() for r in rows(wb, "Applications") if r["Status"] != "To apply"}
        jobs = json.loads((ROOT / "jobs.json").read_text(encoding="utf-8")) if (ROOT / "jobs.json").exists() else []
        # roles at your level first, senior roles after
        self.jobs = sorted((j for j in jobs if j["company"].lower() not in applied),
                           key=lambda j: "senior" in j["title"].lower())
        self.next_job = 0
        self.next_pattern = 0
        self.carry = dict(carry)  # missed units still to place this week
        self.today = today
        self.sunday = today + timedelta(days=6 - today.weekday())

    def take(self, kind, n, day):
        """Add carried-over units to the first block of this kind after today, within this week."""
        extra = self.carry.pop(kind, 0) if self.today < day <= self.sunday else 0
        note = [f"Includes {extra} carried over from earlier this week."] if extra else []
        return n + extra, note

    def apply(self, n, day):
        n, lines = self.take("apply", n, day)
        picks = self.jobs[self.next_job:self.next_job + n]
        self.next_job += n
        lines += [f"- {j['title']} at {j['company']}: {j['url']}" for j in picks]
        if len(picks) < n:
            lines.append(f"- {n - len(picks)} more from the Targets tab (priority A) or Applications 'To apply'")
        lines.append("After each one: log it in Applications, then connect with 2 people there.")
        return f"{n} application{'s' * (n != 1)}", lines

    def leetcode(self, n, day):
        n, lines = self.take("leetcode", n, day)
        picks = [PATTERNS[(self.next_pattern + i) % len(PATTERNS)] for i in range(n)]
        self.next_pattern += n
        return f"{n} LeetCode", lines + [f"- Pattern: {p} (20-minute rule)" for p in picks] + ["Log each in the Learning tab."]

    def outreach(self, n, day):
        n, lines = self.take("outreach", n or 15, day)
        due = [c for c in self.contacts if as_date(c.get("Next date")) and as_date(c["Next date"]) <= day
               and c["Status"] not in ("Referred (written)", "Referred (call)", "Not a fit")]
        find = [c for c in self.contacts if c["Status"] == "Find person"]
        lines += ["Follow-ups due:"] + [f"- {c['Person'] or '?'} ({c['Company']}): {c['Next action']}" for c in due] \
            if due else ["No follow-ups due (check Contacts for new replies)."]
        lines += ["", f"{n} new connection requests with a note (Messages tab). Start with:"]
        lines += [f"- Find someone at {c['Company']} ({c['Job / role']})" for c in find[:5]]
        lines.append("- Then priority A companies in Targets. Log each person in Contacts. Max ~75 a week.")
        return f"{n} connection requests + follow-ups", lines

    def catchup(self, n, day):
        left = {k: v for k, v in self.carry.items() if v} if day == self.sunday else {}
        self.carry = {} if left else self.carry
        if not left:
            return "Catch-up (nothing missed)", ["Nothing to catch up. Use the time for portfolio or rest."]
        return "Catch-up", [f"- {v} {NAMES[k]}" for k, v in left.items()] + ["Finish what you can; the rest resets Monday."]

    def expand(self, title, day):
        notes = []

        def sub(m):
            label, lines = getattr(self, m.group(1))(int(m.group(2) or 0), day)
            notes.extend(lines)
            return label

        return re.sub(r"\{(\w+)(?::(\d+))?\}", sub, title), notes


def fit(blocks, busy):
    """Move task blocks off in-person events. Returns (blocks, log lines)."""
    if not busy:
        return blocks, []
    clash = lambda a, b, spans: next((t for x, y, t in spans if a < y and x < b), None)
    soft = lambda blk: blk[2].startswith(SOFT)
    cur = [list(blk) for blk in blocks]
    gone, log = set(), []
    for i, (s, e, title, *extra) in enumerate(blocks):
        a, b = span(s, e)
        hit = clash(a, b, busy)
        if not hit or "free" in extra or title.startswith(FIXED):
            continue
        if title.startswith(DROP) or soft(cur[i]):
            gone.add(i)
            log.append(f"dropped {title} ({hit})")
            continue
        # tasks may take the time of soft routine blocks (walks, dinner prep), never of real events or other tasks
        occupied = busy + [(*span(x[0], x[1]), x[2]) for j, x in enumerate(cur)
                           if j != i and j not in gone and "free" not in x[3:] and not soft(x)]
        slot = next((m for m in [*range(a, DAY_END - (b - a) + 1, 15), *range(DAY_START, a, 15)]
                     if not clash(m, m + b - a, occupied)), None)
        if slot is None:
            gone.add(i)
            log.append(f"no free slot for {title} ({hit})")
            continue
        cur[i] = [hhmm(slot), hhmm(slot + b - a), title, f"Moved from {s}: clashes with {hit}.", *extra]
        log.append(f"moved {title} {s} -> {hhmm(slot)} ({hit})")
        for j, x in enumerate(cur):
            if j not in gone and soft(x) and clash(*span(x[0], x[1]), [(slot, slot + b - a, title)]):
                gone.add(j)
                log.append(f"dropped {x[2]} (made room for {title})")
    return sorted((x for j, x in enumerate(cur) if j not in gone), key=lambda x: x[0]), log


def build(cfg, start, days, busy, online, carry):
    p = Planner(cfg, carry, start)
    tz = cfg["tz"]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    out = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//santisky8//job-assistant//EN",
           "X-WR-CALNAME:Daily plan", f"X-WR-TIMEZONE:{tz}", "REFRESH-INTERVAL;VALUE=DURATION:PT6H"]
    for i in range(days):
        day = start + timedelta(days=i)
        seen = {}
        blocks, log = fit(cfg["days"][DAYS[day.weekday()]], busy.get(day, []))
        for line in log:
            print(f"  {day}: {line}")
        for n, (s, e, title, *extra) in enumerate(blocks):
            summary, notes = p.expand(title, day)
            a, b = span(s, e)
            notes = [f"Online at the same time: {t}. Watch it while you do this." for x, y, t in online.get(day, [])
                     if a < y and x < b and "free" not in extra] + notes
            # date + SUMMARY is the task key in Lyfe Time, so keep it unique within a day
            seen[summary] = seen.get(summary, 0) + 1
            if seen[summary] > 1:
                summary += f" ({seen[summary]})"
            # "free" marks a reminder that overlaps other blocks without making you busy
            free = "free" in extra
            notes = [x for x in extra if x != "free"] + notes
            d = day.strftime("%Y%m%d")
            end = (day + timedelta(days=1) if mins(e) <= mins(s) else day).strftime("%Y%m%d")
            out += ["BEGIN:VEVENT", f"UID:{d}-{n}@job-assistant", f"DTSTAMP:{stamp}",
                    f"DTSTART;TZID={tz}:{d}T{s.replace(':', '')}00", f"DTEND;TZID={tz}:{end}T{e.replace(':', '')}00",
                    f"SUMMARY:{esc(summary)}", f"TRANSP:{'TRANSPARENT' if free else 'OPAQUE'}"]
            if notes:
                out.append(f"DESCRIPTION:{esc(chr(10).join(notes))}")
            out += ["BEGIN:VALARM", "ACTION:DISPLAY", f"DESCRIPTION:{esc(summary)}", "TRIGGER:-PT5M", "END:VALARM",
                    "END:VEVENT"]
    out.append("END:VCALENDAR")
    return "\r\n".join(out) + "\r\n"


# ---------- metrics (tracker "Metrics" tab) ----------

def metrics(cfg, progress, today):
    tracked = [p for p in progress if p["kind"]]
    monday = today - timedelta(days=today.weekday())
    sheet = [["Task", "This week: done / planned", "Last 4 weeks: done / planned", "4-week rate"]]
    for kind, name in NAMES.items():
        week = [p for p in tracked if p["kind"] == kind and p["day"] >= monday]
        month = [p for p in tracked if p["kind"] == kind]
        d4, n4 = sum(p["done"] for p in month), sum(p["n"] for p in month)
        sheet.append([name.capitalize(), f"{sum(p['done'] for p in week)} / {sum(p['n'] for p in week)}",
                      f"{d4} / {n4}", f"{d4 / n4:.0%}" if n4 else "no data yet"])

    sheet += [[], ["When you actually get tasks done (last 4 weeks)", "Units done", "Blocks checked"]]
    by_hour = defaultdict(lambda: [0, 0])
    for p in tracked:
        if p["at"] and p["done"]:
            by_hour[int(p["at"][:2])][0] += p["done"]
            by_hour[int(p["at"][:2])][1] += 1
    sheet += [[f"{h:02d}:00-{h + 1:02d}:00", *by_hour[h]] for h in sorted(by_hour)] or [["no data yet"]]

    wb = load_workbook(cfg["tracker"], data_only=True)
    apps = [r for r in rows(wb, "Applications") if as_date(r.get("Date applied"))]
    sheet += [[], ["Applications by weekday (all time)", "Sent", "Got a response", "Response rate"]]
    for i, name in enumerate(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]):
        sent = [r for r in apps if as_date(r["Date applied"]).weekday() == i]
        hit = [r for r in sent if r["Status"] in ("Assessment", "Interview", "Offer")]
        sheet.append([name, len(sent), len(hit), f"{len(hit) / len(sent):.0%}" if sent else "-"])

    nights = []
    for p in progress:
        if p["summary"].startswith("Sleep") and p["bed"] and p["wake"]:
            hours = (mins(p["wake"]) - mins(p["bed"])) % (24 * 60) / 60
            nights.append((p["day"], p["bed"], p["wake"], hours))
    sheet += [[], ["Sleep (last 4 weeks, goal 22:00-07:00, 9 h)", "Bed", "Wake", "Hours"]]
    sheet += [[str(d), bed, wake, round(h, 1)] for d, bed, wake, h in sorted(nights)] or [["no data yet"]]
    if nights:
        sheet.append(["Average / nights with 9 h+", "", f"{sum(n[3] for n in nights) / len(nights):.1f} h",
                      f"{sum(n[3] >= 9 for n in nights)} of {len(nights)}"])
    sheet += [[], [f"Updated {datetime.now():%Y-%m-%d %H:%M} from Lyfe Time check-offs. Dates applied come from Applications."]]

    wb = load_workbook(cfg["tracker"])  # keep formulas in the other tabs
    if "Metrics" in wb.sheetnames:
        del wb["Metrics"]
    ws = wb.create_sheet("Metrics", 1)
    head, white = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")  # AA contrast
    for i, r in enumerate(sheet):
        ws.append(r)
        if i == 0 or (i and not sheet[i - 1]):  # first row of each section is its header
            for c in ws[ws.max_row]:
                c.fill, c.font = head, white
    for col, width in zip("ABCD", (44, 26, 28, 16)):
        ws.column_dimensions[col].width = width
    try:
        wb.save(cfg["tracker"])
        print("Updated Metrics tab")
    except PermissionError:
        print("Tracker is open in another program; Metrics tab not updated this time")


# ---------- publish ----------

def publish(cfg, ics):
    if not cfg.get("gist_id"):
        raise SystemExit("Set gist_id in schedule.json first.")
    repo = OUT / "gist"
    if not repo.exists():
        subprocess.run(["git", "clone", "-q", f"https://gist.github.com/{cfg['gist_id']}.git", str(repo)], check=True)
    subprocess.run(["git", "-C", str(repo), "pull", "-q"], check=True)
    (repo / "plan.ics").write_text(ics, encoding="utf-8", newline="")
    subprocess.run(["git", "-C", str(repo), "add", "plan.ics"], check=True)
    if subprocess.run(["git", "-C", str(repo), "diff", "--cached", "--quiet"]).returncode:
        subprocess.run(["git", "-C", str(repo), "commit", "-qm", f"plan {date.today()}"], check=True)
        subprocess.run(["git", "-C", str(repo), "push", "-q"], check=True)
    print(f"Published. Subscribe URL: https://gist.githubusercontent.com/{cfg['github_user']}/{cfg['gist_id']}/raw/plan.ics")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--publish", action="store_true")
    ap.add_argument("--days", type=int, default=14)
    args = ap.parse_args()
    cfg = json.loads((ROOT / "schedule.json").read_text(encoding="utf-8"))
    today = date.today()
    busy, online, progress = read_calendar(cfg, today - timedelta(days=28), today + timedelta(days=args.days))
    carry = missed_this_week(progress, today)
    if carry:
        print("  carried over:", ", ".join(f"{v} {NAMES[k]}" for k, v in carry.items() if v))
    ics = build(cfg, today, args.days, busy, online, carry)
    OUT.mkdir(exist_ok=True)
    (OUT / "plan.ics").write_text(ics, encoding="utf-8", newline="")
    print(f"Wrote {OUT / 'plan.ics'}")
    metrics(cfg, [p for p in progress if p["day"] <= today], today)
    if args.publish:
        publish(cfg, ics)


if __name__ == "__main__":
    main()
